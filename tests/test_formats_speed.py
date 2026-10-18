"""Any-format input (converted to PDF/PNG first) and the two-phase loading of the web app."""
import importlib
import io
import os
import time

import pymupdf
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from conftest import image_with_text, needs_tesseract
from processing import formats, pipeline


def _md(path, name=None):
    return pipeline.process(str(path), name or os.path.basename(str(path)), "t", None, pipeline.Progress())["_markdown"]


@pytest.fixture(autouse=True)
def _builtin_office(monkeypatch):
    monkeypatch.setenv("PARSEFLOW_OFFICE", "builtin")      # deterministic without LibreOffice


def test_word_document(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("Quarterly Report", 1)
    d.add_paragraph("Revenue grew 12% to $4.5M. Café résumé.")
    t = d.add_table(rows=3, cols=3)
    for i, r in enumerate([["Region", "Q1", "Q2"], ["North", "1.2", "1.5"], ["South", "0.8", "0.9"]]):
        for j, v in enumerate(r):
            t.cell(i, j).text = v
    d.add_paragraph("First point", style="List Bullet")
    p = tmp_path / "report.docx"
    d.save(p)
    md = _md(p)
    assert "# Quarterly Report" in md and "Café résumé" in md and "| North | 1.2 | 1.5 |" in md and "- First point" in md
    assert "| Revenue" not in md                                   # the paragraph above is not made a table header


def test_spreadsheet_and_csv(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    for r in (["Product", "Units", "Price"], ["Pen", 10, 1.5], ["Book", 3, 12.0]):
        ws.append(r)
    wb.save(tmp_path / "s.xlsx")
    md = _md(tmp_path / "s.xlsx")
    assert "## Sales" in md and "| Pen | 10 | 1.5 |" in md
    (tmp_path / "people.csv").write_text("name;age\nAsha;21\nRavi;22\n")
    assert "| Asha | 21 |" in _md(tmp_path / "people.csv")


def test_slides_text_html_markdown_json(tmp_path):
    import pptx
    pr = pptx.Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[1])
    s.shapes.title.text = "Project Plan"
    s.placeholders[1].text = "Phase 1: research\nPhase 2: build"
    pr.slides.add_slide(pr.slide_layouts[1]).shapes.title.text = "Budget"
    pr.save(tmp_path / "plan.pptx")
    d = pipeline.process(str(tmp_path / "plan.pptx"), "plan.pptx", "t", None, pipeline.Progress())
    assert len(d["pages"]) == 2 and "# Project Plan" in d["_markdown"] and "- Phase 2: build" in d["_markdown"]
    assert d["converted"] == dict(from_extension=".pptx", method="builtin")
    (tmp_path / "page.html").write_text("<h1>Hello</h1><script>alert(1)</script><p>World <img src='http://x/y.png'></p>")
    md = _md(tmp_path / "page.html")
    assert "# Hello" in md and "World" in md and "alert" not in md
    (tmp_path / "n.md").write_text("# Title\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
    assert "| 1 | 2 |" in _md(tmp_path / "n.md")
    (tmp_path / "u16.txt").write_bytes("Grüße aus Köln\n\nZweiter Absatz".encode("utf-16"))
    assert "Grüße aus Köln" in _md(tmp_path / "u16.txt")
    (tmp_path / "d.json").write_text('{"name": "Asha", "items": [1, 2]}')
    assert '"name": "Asha"' in _md(tmp_path / "d.json")


@needs_tesseract()
def test_multipage_tiff_and_webp(tmp_path):
    a = image_with_text(["Scanned page number one"], size=(1700, 600))
    b = image_with_text(["Second scanned page here"], size=(1700, 600))
    a.save(tmp_path / "scan.tiff", save_all=True, append_images=[b])
    d = pipeline.process(str(tmp_path / "scan.tiff"), "scan.tiff", "t", None, pipeline.Progress())
    assert len(d["pages"]) == 2 and "Second scanned page" in d["_markdown"]
    a.save(tmp_path / "one.webp")
    assert "Scanned page number one" in _md(tmp_path / "one.webp")


def test_unsupported_and_libreoffice_only_formats(tmp_path, monkeypatch):
    (tmp_path / "x.exe").write_bytes(b"MZ")
    with pytest.raises(pipeline.PFError) as e:
        _md(tmp_path / "x.exe")
    assert "not supported" in str(e.value)
    monkeypatch.setattr(formats, "libreoffice", lambda: None)
    (tmp_path / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(pipeline.PFError) as e:
        _md(tmp_path / "old.doc")
    assert "LibreOffice" in str(e.value)


@pytest.mark.skipif(not formats.libreoffice(), reason="LibreOffice not installed")
def test_libreoffice_route(tmp_path, monkeypatch):
    import docx
    monkeypatch.setenv("PARSEFLOW_OFFICE", "libreoffice")
    d = docx.Document()
    d.add_paragraph("Converted by LibreOffice with its real layout.")
    d.save(tmp_path / "lo.docx")
    r = pipeline.process(str(tmp_path / "lo.docx"), "lo.docx", "t", None, pipeline.Progress())
    assert r["converted"]["method"] == "libreoffice" and "real layout" in r["_markdown"]


def _figure_pdf(path):
    buf = io.BytesIO()
    image_with_text(["Throughput 42 units", "Latency 7 ms"], size=(1200, 500), font_px=48).save(buf, "PNG")
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 72), "The chart below shows the measured values.", fontsize=11)
    p.insert_image(pymupdf.Rect(72, 100, 520, 290), stream=buf.getvalue())
    doc.save(str(path))
    return str(path)


@needs_tesseract()
def test_document_opens_before_images_are_read(tmp_path):
    pdf = _figure_pdf(tmp_path / "fig.pdf")
    out = tmp_path / "job"
    d = pipeline.process(pdf, "fig.pdf", "t", str(out), pipeline.Progress(str(out / "status.json")), defer_visuals=True)
    pend = [b for b in d["blocks"] if b.get("pending")]
    assert d["pending_visuals"] == 1 and pend and pend[0]["extractor"] == "pending" and "measured values" in d["_markdown"]
    prog = pipeline.Progress(str(out / "status.json"))
    assert pipeline.enrich_visuals(pdf, "fig.pdf", str(out), prog) == 1
    import json
    doc = json.loads((out / "document.json").read_text())
    fig = [b for b in doc["blocks"] if b["type"] == "figure"]
    assert not any(b.get("pending") for b in doc["blocks"]) and doc["pending_visuals"] == 0
    assert fig and "Throughput" in (fig[0].get("content") or "") and fig[0]["block_id"] == pend[0]["block_id"]
    assert "Throughput" in (out / "document.md").read_text() and prog.enrich["state"] == "done"


@needs_tesseract()
def test_app_upload_any_format_and_background_reading(tmp_path, monkeypatch):
    import app as A
    A = importlib.reload(A)
    c = TestClient(A.app)
    import docx
    d = docx.Document()
    d.add_paragraph("A Word file uploaded to the web app.")
    d.save(tmp_path / "w.docx")
    with open(tmp_path / "w.docx", "rb") as f:
        r = c.post("/api/process", files={"file": ("w.docx", f, "application/octet-stream")}).json()
    assert r["type"] == "DOCX" and r["pages"] == 1
    t0 = time.time()
    while c.get(f"/api/status/{r['document_id']}").json()["status"] == "running":
        assert time.time() - t0 < 60
        time.sleep(0.2)
    assert "A Word file uploaded" in c.get(f"/api/document/{r['document_id']}/markdown").text
    j = c.get(f"/api/page/{r['document_id']}/1.jpg")
    assert j.status_code == 200 and j.headers["content-type"] == "image/jpeg" and Image.open(io.BytesIO(j.content)).size[0] > 600
    thumb = c.get(f"/api/page/{r['document_id']}/1.jpg?w=200")
    assert Image.open(io.BytesIO(thumb.content)).size[0] == 200
    bad = c.post("/api/process", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert bad.status_code == 400 and "DOCX" in bad.json()["error"]["message"]
    # a PDF with a picture: the document is ready first, the picture text arrives in the background
    with open(_figure_pdf(tmp_path / "fig.pdf"), "rb") as f:
        r = c.post("/api/process", files={"file": ("fig.pdf", f, "application/pdf")}).json()
    t0 = time.time()
    while True:
        st = c.get(f"/api/status/{r['document_id']}").json()
        if (st.get("enrich") or {}).get("state") == "done":
            break
        assert time.time() - t0 < 60
        time.sleep(0.2)
    assert "Throughput" in c.get(f"/api/document/{r['document_id']}/markdown").text


def _scan_pdf(path):
    from conftest import scanned_pdf
    return scanned_pdf(path, image_with_text(["A scanned page of text"], size=(1200, 900)))


def test_without_ocr_a_connected_ai_reads_scanned_pages(tmp_path, monkeypatch):
    from processing import ocr, vision
    monkeypatch.setattr(ocr, "available", lambda: None)
    monkeypatch.setattr(vision, "available", lambda: True)
    monkeypatch.setattr(vision, "analyze_crop", lambda img, timeout=60: (
        [dict(type="paragraph", content="A scanned page of text, read by the AI.", confidence=0.9)], None))
    pdf = _scan_pdf(tmp_path / "scan.pdf")
    d = pipeline.process(pdf, "scan.pdf", "t", None, pipeline.Progress())
    assert "read by the AI" in d["_markdown"] and not d["warnings"]
    # web app: the page opens at once and the AI reads it in the background
    out = tmp_path / "job"
    d = pipeline.process(pdf, "scan.pdf", "t", str(out), pipeline.Progress(str(out / "status.json")), defer_visuals=True)
    assert d["pending_visuals"] == 1
    pipeline.enrich_visuals(pdf, "scan.pdf", str(out))
    assert "read by the AI" in (out / "document.md").read_text()


def test_without_ocr_or_ai_one_friendly_note(tmp_path, monkeypatch):
    from processing import ocr, vision
    monkeypatch.setattr(ocr, "available", lambda: None)
    monkeypatch.setattr(vision, "available", lambda: False)
    d = pipeline.process(_scan_pdf(tmp_path / "scan.pdf"), "scan.pdf", "t", None, pipeline.Progress())
    assert len(d["warnings"]) == 1 and "run.bat" in d["warnings"][0] and "not installed" not in d["warnings"][0]
    assert not any("Vision API" in w or "OCR engine unavailable" in w for w in d["warnings"])


def test_launcher_installs_tesseract_once_with_winget(monkeypatch, tmp_path):
    import shutil
    import subprocess
    import start
    from processing import ocr
    calls = []
    monkeypatch.setattr(start, "HERE", tmp_path / "backend")
    monkeypatch.delenv("PARSEFLOW_NO_AUTO_INSTALL", raising=False)
    monkeypatch.setattr(ocr, "available", lambda: None)
    monkeypatch.setattr(start.os, "name", "nt")
    monkeypatch.setattr(shutil, "which", lambda n: "C:/winget.exe" if n == "winget" else None)
    monkeypatch.setattr(subprocess, "run", lambda cmd, **k: calls.append(cmd))
    assert start.ensure_tesseract() is False
    assert calls and calls[0][:4] == ["winget", "install", "--id", "UB-Mannheim.TesseractOCR"]
    calls.clear()
    start.ensure_tesseract()                      # tried recently: no second attempt on the next start
    assert calls == []
    monkeypatch.setattr(ocr, "available", lambda: "tesseract 5")
    assert start.ensure_tesseract() is True
