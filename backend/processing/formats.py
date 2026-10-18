"""Open (almost) any document: everything that is not already a PDF or a PNG/JPEG is turned into a
PDF first, so the one pipeline (text layer, tables, OCR, page viewer, provenance) handles it.

Routes, fastest that keeps the content:
  * PDF, PNG, JPEG ............................ used as they are
  * TIFF (multi-page), BMP, GIF, WebP, ... ... converted to PNG / a PDF of scanned pages, then OCR
  * XPS, OXPS, EPUB, MOBI, FB2, CBZ, SVG ...... MuPDF converts them to PDF (real text layer)
  * XLSX ........................................ built-in: every sheet as a table of its cell values
  * DOCX, PPTX, ODT, ODS, ODP, DOC, XLS, PPT,
    RTF ......................................... LibreOffice (headless) when it is installed: true
                                                  layout and pictures. Without LibreOffice, DOCX,
                                                  PPTX and RTF use the built-in route below.
  * TXT, MD, CSV, TSV, HTML, JSON, XML, LOG ... built-in: HTML laid out to PDF pages by MuPDF

Untrusted input: HTML is laid out without scripts or network access (MuPDF only reads the archive we
give it); LibreOffice runs headless with a private profile and a time limit.
"""
import csv
import html as H
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf

NATIVE = {".pdf", ".png", ".jpg", ".jpeg"}
IMAGES = {".tif", ".tiff", ".bmp", ".gif", ".webp", ".jfif", ".pbm", ".pgm", ".ppm", ".pnm"}
MUPDF = {".xps", ".oxps", ".epub", ".mobi", ".fb2", ".cbz", ".svg"}
TEXTLIKE = {".txt", ".text", ".md", ".markdown", ".csv", ".tsv", ".html", ".htm", ".xhtml", ".json", ".xml", ".log", ".rst", ".tex"}
OFFICE_BUILTIN = {".docx", ".xlsx", ".xlsm", ".pptx", ".rtf"}
OFFICE_LO = {".doc", ".xls", ".ppt", ".odt", ".ods", ".odp", ".docm", ".pptm", ".dotx", ".potx", ".xltx"}
SUPPORTED = NATIVE | IMAGES | MUPDF | TEXTLIKE | OFFICE_BUILTIN | OFFICE_LO
MAX_TABLE_ROWS = int(os.getenv("PARSEFLOW_MAX_TABLE_ROWS", "5000"))

CSS = """body{font-family:sans-serif;font-size:10.5pt;line-height:1.35}
h1{font-size:18pt}h2{font-size:15pt}h3{font-size:13pt}h4,h5,h6{font-size:11.5pt}
table{border-collapse:collapse;margin:6pt 0}td,th{border:0.6pt solid #888;padding:2pt 4pt;vertical-align:top}
th{font-weight:bold}pre,code{font-family:monospace;font-size:9pt}pre{white-space:pre-wrap}
img{max-width:100%}"""


class FormatError(Exception):
    pass


def describe():
    """For /api/health and the UI: what can be opened here."""
    return dict(extensions=sorted(SUPPORTED if libreoffice() else SUPPORTED - OFFICE_LO), libreoffice=bool(libreoffice()))


def libreoffice():
    exe = os.getenv("LIBREOFFICE") or shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        return exe
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.isfile(c):
            return c
    return None


def needs_conversion(filename):
    return Path(filename).suffix.lower() not in NATIVE


def to_pipeline_input(src, filename, out_dir):
    """Convert `src` (uploaded as `filename`) into something the pipeline reads. Returns
    (path, method). The result is written into out_dir as source.pdf or source.png."""
    ext = Path(filename).suffix.lower()
    out_dir = Path(out_dir)
    if ext in NATIVE:
        return str(src), "native"
    if ext not in SUPPORTED:
        raise FormatError(f"'{ext or filename}' files are not supported. ParseFlow opens: " + ", ".join(sorted(e[1:] for e in SUPPORTED)) + ".")
    pdf = out_dir / "source.pdf"
    if ext in (".docx", ".docm", ".xlsx", ".xlsm", ".pptx", ".pptm", ".dotx", ".potx", ".xltx", ".odt", ".ods", ".odp", ".epub"):
        import zipfile
        if not zipfile.is_zipfile(src):        # these formats are zip packages
            raise FormatError(f"This {ext} file is damaged or not really a {ext} file.")
    if ext in IMAGES:
        return _image(src, out_dir), "image"
    if ext in MUPDF:
        with pymupdf.open(str(src)) as d:
            data = d.convert_to_pdf()
        pdf.write_bytes(data)
        return str(pdf), "mupdf"
    if ext in TEXTLIKE:
        _html_to_pdf(_textlike_html(src, ext), pdf)
        return str(pdf), "builtin"
    if ext in (".xlsx", ".xlsm"):
        # spreadsheets: the cell values themselves give exact tables (a printed sheet has no grid lines)
        html, archive, per_page = _office_html(src, ext)
        _html_to_pdf(html, pdf, archive=archive, pages=per_page)
        return str(pdf), "builtin"
    lo = libreoffice()
    office_mode = os.getenv("PARSEFLOW_OFFICE", "builtin" if (os.getenv("RENDER") or os.getenv("PARSEFLOW_LOW_MEM")) else "libreoffice")
    if lo and office_mode != "builtin":
        try:
            _libreoffice(lo, src, ext, pdf)
            return str(pdf), "libreoffice"
        except FormatError:
            if ext not in OFFICE_BUILTIN:
                raise
    if ext in OFFICE_BUILTIN:
        html, archive, per_page = _office_html(src, ext)
        _html_to_pdf(html, pdf, archive=archive, pages=per_page)
        return str(pdf), "builtin"
    raise FormatError(f"Opening {ext} files needs LibreOffice (free: libreoffice.org). Install it, or save the file as "
                      "PDF or DOCX/XLSX/PPTX and upload that.")


# ------------------------------------------------------------------------------------------ images

def _image(src, out_dir):
    from PIL import Image, ImageSequence
    try:
        im = Image.open(src)
        frames = [f.convert("RGB") for f in ImageSequence.Iterator(im)]
    except Exception as e:
        raise FormatError(f"The image could not be read ({type(e).__name__}).")
    if not frames:
        raise FormatError("The image has no frames.")
    if len(frames) == 1 or os.path.splitext(str(src))[1].lower() == ".gif":
        p = Path(out_dir) / "source.png"
        frames[0].save(p, "PNG")
        return str(p)
    p = Path(out_dir) / "source.pdf"           # multi-page TIFF: one scanned page per frame
    try:
        dpi = float((im.info.get("dpi") or (0, 0))[0] or 0)
    except Exception:
        dpi = 0.0
    if not 70 <= dpi <= 1200:      # missing or nonsense resolution: assume a 300 dpi scan
        dpi = 300.0
    frames[0].save(p, "PDF", save_all=True, append_images=frames[1:], resolution=float(dpi))
    return str(p)


# ------------------------------------------------------------------------------------ LibreOffice

def _libreoffice(exe, src, ext, pdf):
    with tempfile.TemporaryDirectory() as td:
        inp = Path(td) / ("input" + ext)
        shutil.copyfile(src, inp)
        profile = Path(td) / "profile"
        cmd = [exe, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore", "--nolockcheck",
               "--convert-to", "pdf", "--outdir", td, str(inp)]
        try:
            subprocess.run(cmd, capture_output=True, timeout=float(os.getenv("LIBREOFFICE_TIMEOUT_S", "120")), check=False)
        except subprocess.TimeoutExpired:
            raise FormatError("LibreOffice took too long to convert this file.")
        out = Path(td) / "input.pdf"
        if not out.exists() or out.stat().st_size == 0:
            raise FormatError("LibreOffice could not convert this file (it may be damaged or password-protected).")
        shutil.copyfile(out, pdf)


# ------------------------------------------------------------------------------- HTML -> PDF pages

def _html_to_pdf(html, pdf, archive=None, pages=None):
    """Lay out HTML on A4 pages. pages: optional list of HTML strings, each starting a new page
    (slides)."""
    parts = pages or [html]
    w = pymupdf.DocumentWriter(str(pdf))
    mb = pymupdf.paper_rect("a4")
    where = mb + (50, 50, -50, -50)
    n = 0
    try:
        for part in parts:
            story = pymupdf.Story(html=part, user_css=CSS, archive=archive)
            more = 1
            while more:
                dev = w.begin_page(mb)
                more, _ = story.place(where)
                story.draw(dev)
                w.end_page()
                n += 1
                if n > 5000:
                    raise FormatError("The document is too long to lay out (over 5,000 pages).")
    finally:
        w.close()
    if n == 0:
        raise FormatError("The document is empty.")


def _read_text(src):
    raw = Path(src).read_bytes()
    for enc in ("utf-8-sig", "utf-16"):
        try:
            if enc == "utf-16" and not raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    try:
        from charset_normalizer import from_bytes
        best = from_bytes(raw).best()
        if best is not None:
            return str(best)
    except Exception:
        pass
    return raw.decode("latin-1")


def _table_html(rows, header=True):
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return ""
    more = len(rows) > MAX_TABLE_ROWS
    rows = rows[:MAX_TABLE_ROWS]
    n = max(len(r) for r in rows)
    out = ["<table>"]
    for i, r in enumerate(rows):
        tag = "th" if header and i == 0 else "td"
        out.append("<tr>" + "".join(f"<{tag}>{H.escape(str(c))}</{tag}>" for c in list(r) + [""] * (n - len(r))) + "</tr>")
    out.append("</table>")
    if more:
        out.append(f"<p><i>(only the first {MAX_TABLE_ROWS} rows are shown)</i></p>")
    return "".join(out)


_UNSAFE = re.compile(r"<(script|style|iframe|object|embed|noscript)\b.*?</\1\s*>", re.S | re.I)


def _textlike_html(src, ext):
    text = _read_text(src)
    if ext in (".csv", ".tsv"):
        dialect = "excel-tab" if ext == ".tsv" else "excel"
        try:
            if ext == ".csv":
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            pass
        return _table_html(list(csv.reader(io.StringIO(text), dialect)))
    if ext in (".md", ".markdown"):
        import markdown as mdlib
        return mdlib.markdown(text, extensions=["tables", "fenced_code"])
    if ext in (".html", ".htm", ".xhtml"):
        body = _UNSAFE.sub("", text)
        body = re.sub(r"<(img|link|source|video|audio)\b[^>]*>", "", body, flags=re.I)   # no external fetches
        return body
    if ext == ".json":
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except Exception:
            pass
        return "<pre>" + H.escape(text) + "</pre>"
    if ext in (".xml", ".log", ".tex"):
        return "<pre>" + H.escape(text) + "</pre>"
    # plain text: blank lines separate paragraphs, single newlines are kept
    paras = re.split(r"\n\s*\n", text.replace("\r\n", "\n"))
    return "".join("<p>" + H.escape(p.strip("\n")).replace("\n", "<br/>") + "</p>" for p in paras if p.strip())


# --------------------------------------------------------------------- built-in office conversion

def _office_html(src, ext):
    if ext == ".docx":
        return _docx_html(src) + (None,)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx_html(src), None, None
    if ext == ".pptx":
        return "", None, _pptx_pages(src)
    if ext == ".rtf":
        return _rtf_html(src), None, None
    raise FormatError(f"No built-in converter for {ext}.")


def _docx_html(src):
    import docx
    from docx.oxml.ns import qn
    try:
        d = docx.Document(src)
    except Exception as e:
        raise FormatError(f"The Word file could not be read ({type(e).__name__}).")
    archive = pymupdf.Archive()
    out, n_img = [], 0
    in_list = False

    def runs_html(p):
        nonlocal n_img
        s = []
        for r in p.runs:
            t = H.escape(r.text)
            if r.bold:
                t = f"<b>{t}</b>"
            if r.italic:
                t = f"<i>{t}</i>"
            s.append(t)
            for blip in r._element.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                part = d.part.related_parts.get(rid) if rid else None
                if part is not None and getattr(part, "blob", None):
                    n_img += 1
                    name = f"img{n_img}" + (os.path.splitext(getattr(part, "partname", "") or "")[1] or ".png")
                    archive.add(part.blob, name)
                    s.append(f'<img src="{name}"/>')
        return "".join(s)

    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(child, d)
            style = (p.style.name if p.style is not None else "") or ""
            inner = runs_html(p)
            m = re.match(r"(Heading|Title)\s*(\d?)", style)
            is_list = "List" in style or child.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None
            if is_list and not in_list:
                out.append("<ul>")
                in_list = True
            if not is_list and in_list:
                out.append("</ul>")
                in_list = False
            if not inner.strip():
                continue
            if m:
                lvl = 1 if m.group(1) == "Title" else min(6, int(m.group(2) or 1))
                out.append(f"<h{lvl}>{inner}</h{lvl}>")
            elif is_list:
                out.append(f"<li>{inner}</li>")
            else:
                out.append(f"<p>{inner}</p>")
        elif tag == "tbl":
            if in_list:
                out.append("</ul>")
                in_list = False
            t = docx.table.Table(child, d)
            rows = []
            for r in t.rows:
                cells, prev = [], None
                for c in r.cells:
                    if prev is not None and c._tc is prev[0]._tc:
                        prev[1] += 1                     # horizontally merged cell
                        continue
                    prev = [c, 1]
                    cells.append(prev)
                rows.append("<tr>" + "".join(f'<td{" colspan=%d" % cs if cs > 1 else ""}>{H.escape(c.text)}</td>' for c, cs in cells) + "</tr>")
            out.append("<table>" + "".join(rows) + "</table>")
    if in_list:
        out.append("</ul>")
    return "".join(out) or "<p></p>", archive


def _xlsx_html(src):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    except Exception as e:
        raise FormatError(f"The spreadsheet could not be read ({type(e).__name__}).")
    out = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else (f"{v:g}" if isinstance(v, float) else str(v)) for v in r])
            if len(rows) > MAX_TABLE_ROWS + 1:
                break
        while rows and not any(c.strip() for c in rows[-1]):
            rows.pop()
        if rows:                                       # drop empty trailing columns
            width = max((max((i + 1 for i, c in enumerate(r) if c.strip()), default=0) for r in rows), default=0)
            rows = [r[:width] for r in rows]
        out.append(f"<h2>{H.escape(ws.title)}</h2>" + (_table_html(rows) or "<p><i>(empty sheet)</i></p>"))
    wb.close()
    return "".join(out)


def _pptx_pages(src):
    import pptx
    try:
        prs = pptx.Presentation(src)
    except Exception as e:
        raise FormatError(f"The presentation could not be read ({type(e).__name__}).")
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        parts, title = [], None
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        parts.append(f"<h1>{H.escape(title or f'Slide {i}')}</h1>")
        shapes = sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
        for sh in shapes:
            if title and sh == slide.shapes.title:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                parts.append(_table_html([[c.text for c in r.cells] for r in sh.table.rows]))
            elif sh.has_text_frame and sh.text_frame.text.strip():
                items = [p for p in sh.text_frame.paragraphs if "".join(r.text for r in p.runs).strip()]
                if len(items) > 1:
                    parts.append("<ul>" + "".join(f"<li>{H.escape(''.join(r.text for r in p.runs))}</li>" for p in items) + "</ul>")
                else:
                    parts.append(f"<p>{H.escape(sh.text_frame.text)}</p>")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append(f"<p><i>Notes: {H.escape(slide.notes_slide.notes_text_frame.text)}</i></p>")
        pages.append("".join(parts))
    return pages or ["<p></p>"]


def _rtf_html(src):
    """Plain text from RTF (no LibreOffice): control words removed, paragraphs kept."""
    t = _read_text(src)
    t = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes.fromhex(m.group(1)).decode("cp1252", "replace"), t)
    t = re.sub(r"\\u(-?\d+)\??", lambda m: chr(int(m.group(1)) % 65536), t)
    t = re.sub(r"{\\\*[^{}]*}", "", t)
    t = re.sub(r"\\(par|line)\b ?", "\n", t)
    t = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", t)
    t = t.replace("{", "").replace("}", "").replace("\\\\", "\\")
    return "".join(f"<p>{H.escape(p.strip())}</p>" for p in re.split(r"\n\s*\n|\n", t) if p.strip())
